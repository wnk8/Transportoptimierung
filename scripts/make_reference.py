"""Generate the reference data from the stored values of the Excel model v11.

Reads ``data/excel/Thema1_Transportkonzept_v11_Beispieldaten.xlsx`` with openpyxl in
read-only mode (``data_only=True``: stored values, no formulas). The workbook is never saved.

Outputs in ``data/reference/``:

* ``expected.json``      key figures (Overview section 1, cockpit, Assumptions C01..C05, status)
* ``sendungen_ref.csv``  sheet Input, rows 14.., columns A..BK, header from row 12
* ``touren_ref.csv``     sheet Touren, rows 14.., columns A..AT, header from row 12
* ``routen_ref.csv``     sheet Routen, rows 14..313, columns A..U, header from row 12
* ``snapshot.json``      sensitivity table (Overview section 6, base case plus 13 scenarios)

Usage: ``uv run python scripts/make_reference.py [path/to/workbook.xlsx]``
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "excel" / "Thema1_Transportkonzept_v11_Beispieldaten.xlsx"
OUT = ROOT / "data" / "reference"

# Cell references (sheet 'Overview & Results', section 1 "Methodenvergleich", rows 10..12)
OVERVIEW_ROWS = {"Ist": 10, "M2": 11, "M3": 12}
OVERVIEW_COLS = {
    "co2_fzgkm_kg": "C",
    "delta_co2_fzgkm": "D",
    "co2_tkm_kg": "E",
    "delta_co2_tkm": "F",
    "kosten_eur": "G",
    "delta_kosten": "H",
    "lkw_km": "I",
    "delta_lkw_km": "J",
    "bahn_sendungen": "K",
    "servicegrad": "L",
}
# Cockpit on sheet Input: parameters D3:D9, key figures J3:J8, N3:N8, R3:R8 (labels in G, K, O)
COCKPIT_PARAMS = {"D3": "A01", "D4": "A02", "D5": "C01", "D6": "C02", "D7": "A05", "D8": "A10", "D9": "A08"}
COCKPIT_VALUE_COLS = (("G", "J"), ("K", "N"), ("O", "R"))
# Assumptions section C1: C01..C05 in D49:D53
ASSUMPTIONS_DERIVED = {"C01": "D49", "C02": "D50", "C03": "D51", "C04": "D52", "C05": "D53"}
# Overview section 7: tours (C193 formed, C194 bundled, C195 shipments in bundled tours)
TOUR_CELLS = {
    "touren_gesamt": "C193",
    "touren_gebuendelt": "C194",
    "sendungen_in_gebuendelten_touren": "C195",
}
# Overview section 5: checks in rows 142..161 (B name, C target, D actual, E diff, F status), total C162
CHECK_ROWS = range(142, 162)
STATUS_CELL = "C162"
# Overview section 6: sensitivity rows 166..179 (B name, C parameter, D..J values)
SNAPSHOT_ROWS = range(166, 180)
SNAPSHOT_COLS = {
    "delta_co2_m2": "D",
    "delta_co2_m3_fzgkm": "E",
    "delta_co2_m3_tkm": "F",
    "delta_kosten_m3": "G",
    "bahn_sendungen": "H",
    "touren_gebuendelt": "I",
    "delta_lkw_km_m3": "J",
}
# Data blocks: (sheet, header row, first data row, last data row, last column)
BLOCKS = {
    "sendungen_ref.csv": ("Input", 12, 14, 2513, "BK"),
    "touren_ref.csv": ("Touren", 12, 14, 2513, "AT"),
    "routen_ref.csv": ("Routen", 12, 14, 313, "U"),
}


def _plain(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def _cell(ws: Any, ref: str) -> Any:
    return _plain(ws[ref].value)


def _header(text: Any) -> str:
    return str(text).split("\n", 1)[0].strip() if text is not None else ""


def build_expected(wb: Any, source_name: str) -> dict[str, Any]:
    ov = wb["Overview & Results"]
    inp = wb["Input"]
    ass = wb["Assumptions"]
    kennzahlen = {
        name: {key: _cell(ov, f"{col}{row}") for key, col in OVERVIEW_COLS.items()}
        for name, row in OVERVIEW_ROWS.items()
    }
    cockpit_params = {pid: _cell(inp, ref) for ref, pid in COCKPIT_PARAMS.items()}
    cockpit_values: dict[str, Any] = {}
    for label_col, value_col in COCKPIT_VALUE_COLS:
        for row in range(3, 9):
            label = inp[f"{label_col}{row}"].value
            if label:
                cockpit_values[str(label)] = _cell(inp, f"{value_col}{row}")
    checks = [
        {
            "nr": i + 1,
            "name": ov[f"B{row}"].value,
            "soll": _cell(ov, f"C{row}"),
            "ist": _cell(ov, f"D{row}"),
            "status": ov[f"F{row}"].value,
        }
        for i, row in enumerate(CHECK_ROWS)
    ]
    return {
        "quelle_datei": source_name,
        "zellbezuege": {
            "kennzahlen": "Overview & Results!C10:L12",
            "cockpit": "Input!D3:R9",
            "abgeleitet": "Assumptions!D49:D53",
            "touren": "Overview & Results!C193:C195",
            "pruefungen": "Overview & Results!B142:F161",
            "gesamtstatus": "Overview & Results!C162",
        },
        "anzahl_sendungen": _cell(ass, "D53"),
        "kennzahlen": kennzahlen,
        "cockpit_parameter": cockpit_params,
        "cockpit": cockpit_values,
        "abgeleitet": {pid: _cell(ass, ref) for pid, ref in ASSUMPTIONS_DERIVED.items()},
        **{key: _cell(ov, ref) for key, ref in TOUR_CELLS.items()},
        "pruefungen": checks,
        "gesamtstatus": _cell(ov, STATUS_CELL),
    }


def build_snapshot(wb: Any) -> dict[str, Any]:
    ov = wb["Overview & Results"]
    rows = []
    for row in SNAPSHOT_ROWS:
        rows.append(
            {
                "zeile": row,
                "szenario": ov[f"B{row}"].value,
                "parameter": ov[f"C{row}"].value,
                **{key: _cell(ov, f"{col}{row}") for key, col in SNAPSHOT_COLS.items()},
            }
        )
    return {"zellbezug": "Overview & Results!B166:J179", "szenarien": rows}


def write_block(wb: Any, name: str) -> int:
    sheet, header_row, first, last, last_col = BLOCKS[name]
    ws = wb[sheet]
    ncol = column_index_from_string(last_col)
    letters = [get_column_letter(i) for i in range(1, ncol + 1)]
    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=ncol, values_only=True))
    header = [f"{letter} {_header(text)}".strip() for letter, text in zip(letters, header_cells, strict=True)]
    count = 0
    with (OUT / name).open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["excel_row", *header])
        for offset, values in enumerate(
            ws.iter_rows(min_row=first, max_row=last, max_col=ncol, values_only=True)
        ):
            plain = [_plain(v) for v in values]
            if all(v is None or v == "" for v in plain[1:]):
                continue
            writer.writerow(
                [
                    first + offset,
                    *("" if v is None else (repr(v) if isinstance(v, float) else v) for v in plain),
                ]
            )
            count += 1
    return count


def compare_with_package(expected: dict[str, Any]) -> list[str]:
    """Compare against section 2.5 of docs/Gesamtpaket.md (rounded values)."""
    doc = {
        "Ist": {
            "co2_fzgkm_kg": 306924.11,
            "co2_tkm_kg": 237300.81,
            "kosten_eur": 841804.20,
            "lkw_km": 812757.62,
        },
        "M2": {
            "co2_fzgkm_kg": 255084.07,
            "co2_tkm_kg": 181826.33,
            "kosten_eur": 631813.87,
            "lkw_km": 583221.01,
        },
        "M3": {
            "co2_fzgkm_kg": 143296.67,
            "co2_tkm_kg": 96792.78,
            "kosten_eur": 631813.87,
            "lkw_km": 262563.03,
            "bahn_sendungen": 551,
        },
    }
    issues = []
    for scen, values in doc.items():
        for key, want in values.items():
            got = expected["kennzahlen"][scen][key]
            if abs(got - want) > 0.006:
                issues.append(f"{scen}.{key}: Excel {got} vs Gesamtpaket {want}")
    for key, want in {"touren_gesamt": 428, "touren_gebuendelt": 318}.items():
        if expected[key] != want:
            issues.append(f"{key}: Excel {expected[key]} vs Gesamtpaket {want}")
    if expected["gesamtstatus"] != "Alle Prüfungen bestanden":
        issues.append(f"gesamtstatus: {expected['gesamtstatus']}")
    return issues


def main() -> int:
    workbook = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    if not workbook.exists():
        print(f"Workbook not found: {workbook}")
        return 2
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        expected = build_expected(wb, workbook.name)
        snapshot = build_snapshot(wb)
        (OUT / "expected.json").write_text(
            json.dumps(expected, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        (OUT / "snapshot.json").write_text(
            json.dumps(snapshot, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        counts = {name: write_block(wb, name) for name in BLOCKS}
    finally:
        wb.close()

    print("Kennzahlen (Overview & Results, Abschnitt 1):")
    for scen, values in expected["kennzahlen"].items():
        print(
            f"  {scen}: CO2e Fzg-km {values['co2_fzgkm_kg']:.2f} kg, tkm {values['co2_tkm_kg']:.2f} kg, "
            f"Kosten {values['kosten_eur']:.2f} EUR, Lkw-km {values['lkw_km']:.2f}, "
            f"Bahn {values['bahn_sendungen']}, Servicegrad {values['servicegrad']:.4f}"
        )
    print(f"  Touren {expected['touren_gesamt']}, gebündelt {expected['touren_gebuendelt']}")
    print(f"  C04 {expected['abgeleitet']['C04']}, Status: {expected['gesamtstatus']}")
    print(f"Zeilen: {counts}, Snapshot-Szenarien: {len(snapshot['szenarien'])}")
    issues = compare_with_package(expected)
    print("Abgleich mit Gesamtpaket 2.5:", "OK" if not issues else "ABWEICHUNG")
    for issue in issues:
        print("  " + issue)
    return 1 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
