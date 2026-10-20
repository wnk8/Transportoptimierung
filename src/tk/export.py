"""Export im Input-Schema (CSV, XLSX) für den Round Trip in die Excel.

Die Ausgabe enthält nur die Eingabespalten A bis H; die Excel rechnet I bis BK selbst.
Das Original wird nie verändert: Der XLSX-Export ist eine neue Arbeitsmappe, der
Round Trip schreibt in eine Kopie der Vorlage.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import shutil
from collections.abc import Sequence
from pathlib import Path

import openpyxl

from tk.excelfn import general_text
from tk.models import INPUT_COLUMNS, INPUT_HEADERS, Shipment


def _cell_text(value: object, column: int) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and column in (3, 5, 6, 7):
        return general_text(value).replace(".", ",")
    return general_text(value)


def to_csv(shipments: Sequence[Shipment]) -> str:
    """CSV mit Semikolon und Dezimalkomma (öffnet in deutschem Excel direkt)."""
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow(INPUT_HEADERS)
    for s in shipments:
        writer.writerow([_cell_text(v, i) for i, v in enumerate(s.cells())])
    return out.getvalue()


def to_xlsx_bytes(shipments: Sequence[Shipment]) -> bytes:
    """Neue Arbeitsmappe mit einem Blatt "Sendungen" (Zeile 1 Kopf, ab Zeile 2 Daten, Spalten A bis H)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sendungen"
    ws.append(list(INPUT_HEADERS))
    for s in shipments:
        ws.append(list(s.cells()))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_into_template(shipments: Sequence[Shipment], template: Path, target: Path) -> Path:
    """Round Trip: Kopie der Vorlage, Sendungen in Input!A14:H eintragen, Original unverändert.

    openpyxl verliert beim Speichern Diagramme (Gesamtpaket 16); deshalb nur auf Kopien.
    Die Kopie muss danach in Excel geöffnet und neu berechnet werden.
    """
    if target.resolve() == template.resolve():
        raise ValueError("Ziel darf nicht die Vorlage sein")
    shutil.copyfile(template, target)
    target.chmod(0o644)
    wb = openpyxl.load_workbook(target)
    try:
        ws = wb["Input"]
        for i, s in enumerate(shipments):
            for j, value in enumerate(s.cells()):
                ws[f"{INPUT_COLUMNS[j]}{14 + i}"] = value
        wb.save(target)
    finally:
        wb.close()
    return target
