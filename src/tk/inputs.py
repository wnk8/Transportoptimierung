"""Import von Sendungen im Schema Input!A:H (Einfügen aus Excel, CSV, XLSX).

Zahlen in Textform (``266,4`` oder ``266.4``) werden wie beim Einfügen in Excel zu Zahlen.
Alles andere bleibt Text und wird von der Eingabeprüfung (Spalte I) bewertet. Eine
Kopfzeile wird erkannt, wenn die erste Zelle kein Datum ist und "Datum" enthält.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import openpyxl

from tk.excelfn import Cell
from tk.models import INPUT_FIELDS, Shipment

_GERMAN_NUMBER = re.compile(r"^[+-]?\d{1,3}(\.\d{3})+(,\d+)?$|^[+-]?\d+,\d+$")
_PLAIN_NUMBER = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
NUMERIC_COLUMNS = {3, 5, 6, 7}  # D Gewicht, F Entfernung, G Paletten, H Fracht


def convert_cell(text: str, column: int) -> Cell:
    """Textzelle in den Wert, den Excel beim Einfügen erzeugt (nur Zahlenspalten D, F, G, H)."""
    value = text.strip()
    if value == "":
        return None
    if column in NUMERIC_COLUMNS:
        if _GERMAN_NUMBER.match(value):
            return float(value.replace(".", "").replace(",", "."))
        if _PLAIN_NUMBER.match(value):
            return float(value)
    return value


def _is_header(cells: list[str]) -> bool:
    return bool(cells) and "datum" in cells[0].strip().lower()


def rows_to_shipments(rows: list[list[str]]) -> list[Shipment]:
    if rows and _is_header(rows[0]):
        rows = rows[1:]
    out = []
    for cells in rows:
        padded = (list(cells) + [""] * len(INPUT_FIELDS))[: len(INPUT_FIELDS)]
        values = [convert_cell(c, i) for i, c in enumerate(padded)]
        ship = Shipment(**dict(zip(INPUT_FIELDS, values, strict=True)))
        if not ship.is_empty():
            out.append(ship)
    return out


def parse_text(text: str) -> list[Shipment]:
    """Eingefügter Text: Tabulator (aus Excel), Semikolon oder Komma als Trenner."""
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []
    sample = lines[0]
    delimiter = "\t" if "\t" in sample else ";" if ";" in sample else ","
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    return rows_to_shipments([list(r) for r in reader])


def parse_csv_bytes(data: bytes) -> list[Shipment]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return parse_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise ValueError("Datei ist weder UTF-8 noch Windows-1252")


def parse_xlsx(path_or_bytes: Path | bytes) -> list[Shipment]:
    """XLSX im Input-Schema: Modellmappe (Input ab Zeile 14) oder erstes Blatt ab Zeile 1 mit Kopf."""
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if "Input" in wb.sheetnames and "Assumptions" in wb.sheetnames:
            ws, first = wb["Input"], 14
        else:
            ws, first = wb.worksheets[0], 1
        out = []
        for i, values in enumerate(ws.iter_rows(min_row=first, max_col=8, values_only=True)):
            cells = list(values) + [None] * (8 - len(values))
            if i == 0 and first == 1 and isinstance(cells[0], str) and "datum" in cells[0].lower():
                continue
            ship = Shipment(**dict(zip(INPUT_FIELDS, cells, strict=True)))
            out.append(ship)
        while out and out[-1].is_empty():
            out.pop()
        return out
    finally:
        wb.close()
