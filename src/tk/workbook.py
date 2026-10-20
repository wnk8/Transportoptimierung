"""Read-only access to the Excel model v11 (stored values, never saved).

The workbook is read once with openpyxl (``read_only=True``, ``data_only=True``) and the
needed sheets are cached as a pickle next to the project data (``data/cache``), keyed by
file name, size and modification time.
"""

from __future__ import annotations

import pickle
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
DEFAULT_WORKBOOK = DATA / "excel" / "Thema1_Transportkonzept_v11_Beispieldaten.xlsx"
CACHE_DIR = DATA / "cache"
CACHE_VERSION = 1

SHEETS = ("Assumptions", "PLZ_Orte")
INPUT_FIRST_ROW = 14
INPUT_LAST_COL = "H"

Grid = dict[tuple[int, int], Any]


@dataclass(slots=True)
class WorkbookData:
    """Cell values of the sheets needed by the core plus the defined names."""

    source: str
    names: dict[str, tuple[str, str]]  # name -> (sheet, range without $)
    sheets: dict[str, Grid]
    input_rows: list[tuple[int, tuple[Any, ...]]]  # (excel row, values A..H)

    def cell(self, sheet: str, ref: str) -> Any:
        col, row = _split_ref(ref)
        return self.sheets[sheet].get((row, col))

    def name_cells(self, name: str) -> list[tuple[str, int, int]]:
        """All cells of a defined name as (sheet, row, col)."""
        sheet, ref = self.names[name]
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        assert min_col is not None and min_row is not None and max_col is not None and max_row is not None
        return [(sheet, r, c) for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]

    def name_values(self, name: str) -> list[Any]:
        return [self.sheets[s].get((r, c)) for s, r, c in self.name_cells(name)]

    def name_value(self, name: str) -> Any:
        values = self.name_values(name)
        if len(values) != 1:
            raise ValueError(f"{name} is not a single cell")
        return values[0]


def _split_ref(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Z]+)\$?(\d+)", ref)
    if not m:
        raise ValueError(ref)
    return column_index_from_string(m.group(1)), int(m.group(2))


def _read(path: Path) -> WorkbookData:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names: dict[str, tuple[str, str]] = {}
        for name, defn in wb.defined_names.items():
            dests = list(defn.destinations)
            if len(dests) == 1 and ":" not in name:
                sheet, ref = dests[0]
                names[name] = (sheet, ref.replace("$", ""))
        sheets: dict[str, Grid] = {}
        for sheet in SHEETS:
            grid: Grid = {}
            for row in wb[sheet].iter_rows():
                for c in row:
                    value = getattr(c, "value", None)
                    if value is not None:
                        grid[(c.row, c.column)] = value
            sheets[sheet] = grid
        last_col = column_index_from_string(INPUT_LAST_COL)
        input_rows = [
            (INPUT_FIRST_ROW + i, tuple(values))
            for i, values in enumerate(
                wb["Input"].iter_rows(min_row=INPUT_FIRST_ROW, max_col=last_col, values_only=True)
            )
            if any(v is not None and v != "" for v in values)
        ]
    finally:
        wb.close()
    return WorkbookData(path.name, names, sheets, input_rows)


def load_workbook_data(path: Path | None = None, use_cache: bool = True) -> WorkbookData:
    """Load the workbook (or its cache). The workbook file is opened read-only."""
    path = path or DEFAULT_WORKBOOK
    if not path.exists():
        raise FileNotFoundError(f"Excel-Mappe nicht gefunden: {path}")
    stat = path.stat()
    cache = CACHE_DIR / f"{path.stem}-{stat.st_size}-{int(stat.st_mtime)}-v{CACHE_VERSION}.pkl"
    # The cache is written only by this function from the local workbook (trusted, never
    # downloaded). It is ignored by git and rebuilt when the workbook changes.
    if use_cache and cache.exists():
        with cache.open("rb") as fh:
            data: WorkbookData = pickle.load(fh)
            return data
    data = _read(path)
    if use_cache:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        with cache.open("wb") as fh:
            pickle.dump(data, fh)
    return data
