"""S16: Export im Input-Schema und Round Trip; Import aus Einfügen, CSV und XLSX."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from tests.conftest import ROOT, needs_excel
from tests.refdata import model_and_result
from tk.export import to_csv, to_xlsx_bytes, write_into_template
from tk.inputs import parse_csv_bytes, parse_text, parse_xlsx
from tk.scenarios import calculate, workbook_shipments
from tk.workbook import load_workbook_data

TEMPLATE = ROOT / "data" / "excel" / "Thema1_Transportkonzept_v11.xlsx"


def test_parse_pasted_excel_rows() -> None:
    text = (
        "Datum\tStart\tZiel\tGewicht\n"
        "2025-01-01\t70173 Stuttgart\t55116 Mainz\t266,4\n"
        "01.02.2025\t68159\t\t1.200\t\t150"
    )
    ships = parse_text(text)
    assert len(ships) == 2
    assert ships[0].gewicht_kg == 266.4
    assert ships[1].gewicht_kg == 1200.0 and ships[1].km == 150.0 and ships[1].ziel is None


def test_text_stays_text_for_validation() -> None:
    ships = parse_text("2025-01-01;70173;55116;viel")
    assert ships[0].gewicht_kg == "viel"


@needs_excel
def test_csv_round_trip_keeps_all_results() -> None:
    model, result = model_and_result()
    ships = workbook_shipments(load_workbook_data())
    again = parse_csv_bytes(to_csv(ships).encode("utf-8-sig"))
    assert len(again) == len(ships)
    re = calculate(again, model)
    for name in ("Ist", "M2", "M3"):
        assert re.kpis[name] == result.kpis[name]


@needs_excel
def test_xlsx_round_trip() -> None:
    model, result = model_and_result()
    ships = workbook_shipments(load_workbook_data())
    again = parse_xlsx(to_xlsx_bytes(ships))
    assert calculate(again, model).kpis["M3"] == result.kpis["M3"]


@needs_excel
def test_write_into_template_copy(tmp_path: Path) -> None:
    """Round Trip in eine Testkopie der Vorlage: Eingaben A:H identisch, Original unverändert."""
    model, result = model_and_result()
    ships = workbook_shipments(load_workbook_data())[:50]
    before = TEMPLATE.stat().st_mtime
    target = write_into_template(ships, TEMPLATE, tmp_path / "roundtrip.xlsx")
    assert TEMPLATE.stat().st_mtime == before
    wb = openpyxl.load_workbook(target, read_only=True)
    rows = list(wb["Input"].iter_rows(min_row=14, max_row=63, max_col=8, values_only=True))
    wb.close()
    assert [tuple(r) for r in rows] == [s.cells() for s in ships]
    reread = parse_xlsx(target)
    assert calculate(reread, model).kpis["Ist"] == calculate(ships, model).kpis["Ist"]
    _ = result
